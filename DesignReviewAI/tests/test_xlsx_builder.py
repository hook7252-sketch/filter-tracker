"""xlsx_builder에 대한 오프라인(비-API) 회귀 테스트.

Claude API 호출 없이, 손으로 만든 샘플 분석 결과로 엑셀 구조와 수식이
기대한 형태로 생성되는지 확인한다.
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from openpyxl import load_workbook

from design_review.schemas import (
    ConsistencyIssue,
    CostRow,
    MissingQuantityRow,
    PageReview,
    ProjectMeta,
    ReviewAnalysis,
    StandardRef,
    StructuralRow,
)
from design_review.xlsx_builder import build_workbook

EXPECTED_SHEETS = ["종합요약", "페이지별검토", "정합성검증", "공사비영향", "구조검토", "기준참고"]


def _sample_meta() -> ProjectMeta:
    return ProjectMeta(
        project_name="테스트 프로젝트",
        source_filename="테스트.pdf",
        review_scope="PDF 3쪽",
        review_date="2026-07-30",
    )


def _sample_reviews() -> list[PageReview]:
    return [
        PageReview(1, "표지", "표지 원문", "식별", "확인", "이상 없음", "없음", "1", ""),
        PageReview(2, "내역서", "앵커 수량", "수량대조", "승인보류", "수량 불일치", "정정 필요", "3", "I01"),
        PageReview(3, "도면", "배치도", "도면검토", "일반보완", "축척 누락", "축척 추가", "1", ""),
    ]


def _sample_analysis() -> ReviewAnalysis:
    return ReviewAnalysis(
        issues=[
            ConsistencyIssue(
                id="I01", importance="승인보류", category="내역", item="앵커 수량 불일치",
                ref1_pages="2", ref1_value="100EA", ref2_pages="3", ref2_value="10EA",
                difference="90EA 과다", impact_verdict="직접비 과다", required_action="정정",
                cost_linkage="확정", related_pages="2,3",
            ),
        ],
        structural_rows=[
            StructuralRow("시설A", "지주", "3", "계산서", "검토결과", "영향", "조치", "재계산", True),
        ],
        cost_rows=[
            CostRow("앵커 설치", 100, 10, "EA", 1000, "확정", True, "근거"),
            CostRow("패널", 50, 60, "m2", 2000, "대안", False, "근거"),
        ],
        missing_quantity_rows=[
            MissingQuantityRow("기초 콘크리트", 1.5, "1m당", 15.0, "총 10m", "m3", "근거"),
        ],
        standards=[
            StandardRef("구조", "KDS 예시 기준", "관련성 설명", "", ""),
        ],
        overall_judgement="종합판정 예시 문구입니다.",
        review_limitation="한계 설명 예시입니다.",
    )


class XlsxBuilderTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.out_path = Path(self.tmpdir.name) / "review.xlsx"
        build_workbook(self.out_path, _sample_meta(), _sample_reviews(), _sample_analysis())
        self.wb = load_workbook(self.out_path)

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_sheet_names_and_order(self):
        self.assertEqual(self.wb.sheetnames, EXPECTED_SHEETS)

    def test_page_review_rows(self):
        ws = self.wb["페이지별검토"]
        self.assertEqual(ws["A5"].value, "PDF쪽")
        self.assertEqual(ws["A6"].value, 1)
        self.assertEqual(ws["E7"].value, "승인보류")  # 2번 페이지 판정

    def test_consistency_issue_row(self):
        ws = self.wb["정합성검증"]
        self.assertEqual(ws["A5"].value, "I01")
        self.assertEqual(ws["D5"].value, "2: 100EA")

    def test_cost_sheet_formulas(self):
        ws = self.wb["공사비영향"]
        self.assertEqual(ws["D9"].value, "=C9-B9")
        self.assertEqual(ws["G9"].value, "=D9*F9")
        self.assertTrue(str(ws["B4"].value).startswith("=SUMIFS("))

    def test_summary_sheet_links_and_counts(self):
        ws = self.wb["종합요약"]
        # 판정 카운트가 COUNTIF 수식으로 생성되는지
        found = any(
            isinstance(ws.cell(row=r, column=2).value, str)
            and ws.cell(row=r, column=2).value.startswith("=COUNTIF(페이지별검토!")
            for r in range(13, 20)
        )
        self.assertTrue(found)
        # 공사비영향 시트로의 링크
        found_link = any(
            ws.cell(row=r, column=4).value == "=공사비영향!B4" for r in range(20, 30)
        )
        self.assertTrue(found_link)


if __name__ == "__main__":
    unittest.main()
