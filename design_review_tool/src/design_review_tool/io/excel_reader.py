from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator, Sequence

import openpyxl
import xlrd


class UnsupportedExcelFormat(ValueError):
    pass


class SheetReader:
    """확장자에 따라 xlrd(.xls)/openpyxl(.xlsx, .xlsm)로 분기해 동일한 인터페이스로 행을 읽는다."""

    def __init__(self, path: str):
        self.path = path
        suffix = Path(path).suffix.lower()
        if suffix == ".xls":
            self._backend: _Backend = _XlrdBackend(path)
        elif suffix in (".xlsx", ".xlsm"):
            self._backend = _OpenpyxlBackend(path)
        else:
            raise UnsupportedExcelFormat(f"지원하지 않는 확장자입니다: {suffix} ({path})")

    @property
    def sheet_names(self) -> list[str]:
        return self._backend.sheet_names

    def rows(self, sheet_name: str) -> Iterator[Sequence[Any]]:
        yield from self._backend.rows(sheet_name)

    def preview_rows(self, sheet_name: str, n: int = 5) -> list[Sequence[Any]]:
        out = []
        for i, row in enumerate(self.rows(sheet_name)):
            if i >= n:
                break
            out.append(row)
        return out


class _Backend:
    @property
    def sheet_names(self) -> list[str]:  # pragma: no cover - interface
        raise NotImplementedError

    def rows(self, sheet_name: str) -> Iterator[Sequence[Any]]:  # pragma: no cover - interface
        raise NotImplementedError


class _XlrdBackend(_Backend):
    def __init__(self, path: str):
        self._wb = xlrd.open_workbook(path)

    @property
    def sheet_names(self) -> list[str]:
        return self._wb.sheet_names()

    def rows(self, sheet_name: str) -> Iterator[Sequence[Any]]:
        sh = self._wb.sheet_by_name(sheet_name)
        for r in range(sh.nrows):
            yield sh.row_values(r)


class _OpenpyxlBackend(_Backend):
    def __init__(self, path: str):
        self._wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    @property
    def sheet_names(self) -> list[str]:
        return self._wb.sheetnames

    def rows(self, sheet_name: str) -> Iterator[Sequence[Any]]:
        ws = self._wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            # xlrd는 빈 셀을 ""로 돌려주므로 동일한 의미가 되도록 None을 ""로 맞춘다.
            yield tuple(v if v is not None else "" for v in row)
